from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(0x0F, 0x12, 0x20)
BG2 = RGBColor(0x18, 0x1C, 0x33)
ACCENT = RGBColor(0xA8, 0xB4, 0xFF)
ORANGE = RGBColor(0xF0, 0xA8, 0x30)
TEAL = RGBColor(0x38, 0xB2, 0xAC)
RED = RGBColor(0xE5, 0x3E, 0x5E)
GREEN = RGBColor(0x4A, 0xDE, 0x80)
WHITE = RGBColor(0xE6, 0xE9, 0xF5)
MUTED = RGBColor(0x8A, 0x93, 0xC4)

def set_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_textbox(slide, left, top, width, height, text, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT, font='Calibri'):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font
    return tb

def add_bullets(slide, left, top, width, height, items, size=16, color=WHITE, bullet_color=None):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_after = Pt(8)
        p.space_before = Pt(4)
        run = p.add_run()
        run.text = "▸  " + item
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = 'Calibri'
    return tb

def add_accent_bar(slide, left, top, height, color=ACCENT, width=0.08):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_rounded_box(slide, left, top, width, height, fill_color, text="", text_size=14, text_color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = fill_color
    if text:
        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.size = Pt(text_size)
        run.font.color.rgb = text_color
        run.font.bold = True
        run.font.name = 'Calibri'
    return shape

# ============================================================
# SLIDE 1: TITLE
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BG)

add_accent_bar(slide, 1.0, 2.2, 2.6, ACCENT, 0.12)
add_textbox(slide, 1.3, 2.2, 11, 1.4, "Distributed Rollback Recovery Simulator", size=44, color=ACCENT, bold=True)
add_textbox(slide, 1.3, 3.5, 11, 0.8, "Visualizing Checkpointing, Consistent Global States & Rollback Recovery", size=20, color=MUTED)

add_rounded_box(slide, 1.3, 4.8, 3.0, 0.6, BG2, "Distributed Computing", 16, ACCENT)

add_textbox(slide, 1.3, 6.0, 11, 0.5, "MCA SEM 1  |  Project TINY", size=14, color=MUTED)

# ============================================================
# SLIDE 2: PROBLEM STATEMENT
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BG)

add_accent_bar(slide, 0.7, 0.6, 0.9, RED, 0.12)
add_textbox(slide, 0.95, 0.55, 11, 0.9, "Problem Statement", size=36, color=RED, bold=True)

add_textbox(slide, 0.7, 1.7, 12, 1.0,
    "In a distributed system, processes run independently and exchange messages. "
    "When a process fails, the system must roll back to a previous checkpoint to recover. "
    "But rolling back processes independently can leave the system in an INCONSISTENT state.",
    size=18, color=WHITE)

add_rounded_box(slide, 0.7, 3.0, 12.0, 1.0, BG2,
    "❝ A message may appear to be RECEIVED but NEVER SENT — called an Orphan Message ❞",
    18, ORANGE)

add_textbox(slide, 0.7, 4.3, 12, 0.5, "Key Questions Addressed:", size=16, color=ACCENT, bold=True)
add_bullets(slide, 0.7, 4.8, 12, 2.5, [
    "How do we find a consistent global state after a failure?",
    "What is the recovery line — the safe point all processes can roll back to?",
    "How does uncoordinated checkpointing cause the Domino Effect?",
    "How does coordinated checkpointing avoid cascading rollbacks?",
], size=16)

# ============================================================
# SLIDE 3: WHAT IS ROLLBACK RECOVERY?
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BG)

add_accent_bar(slide, 0.7, 0.6, 0.9, TEAL, 0.12)
add_textbox(slide, 0.95, 0.55, 11, 0.9, "What is Rollback Recovery?", size=36, color=TEAL, bold=True)

add_bullets(slide, 0.7, 1.8, 12, 5.0, [
    "Checkpoint: A saved snapshot of a process's local state at a point in time.",
    "When a process fails, it restarts from its most recent checkpoint instead of from scratch.",
    "However, processes depend on each other through messages.",
    "Rolling back one process may invalidate messages received by others.",
    "Goal: Find a set of checkpoints (one per process) that forms a CONSISTENT GLOBAL STATE.",
    "A consistent state has NO orphan messages — every received message was sent before the cut.",
], size=17)

# ============================================================
# SLIDE 4: KEY CONCEPTS
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BG)

add_accent_bar(slide, 0.7, 0.6, 0.9, ORANGE, 0.12)
add_textbox(slide, 0.95, 0.55, 11, 0.9, "Key Concepts", size=36, color=ORANGE, bold=True)

concepts = [
    ("Checkpoint", "A saved local state of a process. Shown as orange hexagons.", ORANGE),
    ("Consistent Global State", "A cut through all timelines with no orphan messages.", ACCENT),
    ("Orphan Message", "A message whose receive is in the cut but send is not.", RED),
    ("Recovery Line", "The most recent consistent set of checkpoints to roll back to.", GREEN),
    ("Domino Effect", "Cascading rollbacks across processes due to uncoordinated checkpointing.", TEAL),
]

for i, (title, desc, color) in enumerate(concepts):
    y = 1.7 + i * 1.1
    add_rounded_box(slide, 0.7, y, 0.45, 0.7, color, "", 12)
    add_textbox(slide, 1.3, y, 11.3, 0.4, title, size=16, color=color, bold=True)
    add_textbox(slide, 1.3, y + 0.4, 11.3, 0.6, desc, size=14, color=MUTED)

# ============================================================
# SLIDE 5: THE RECOVERY ALGORITHM
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BG)

add_accent_bar(slide, 0.7, 0.6, 0.9, GREEN, 0.12)
add_textbox(slide, 0.95, 0.55, 11, 0.9, "The Recovery Line Algorithm", size=36, color=GREEN, bold=True)

add_textbox(slide, 0.7, 1.7, 12, 0.5, "Iterative Fixpoint Approach (Orphan-Message Elimination):", size=17, color=ACCENT, bold=True)

steps = [
    ("1", "Initialize: Set each process's frontier to its latest checkpoint."),
    ("2", "Scan all messages. For each message, check if it is an ORPHAN: the receive event is within the cut but the send event is NOT."),
    ("3", "If orphan found: Roll back the RECEIVING process to its checkpoint immediately BEFORE the receive event."),
    ("4", "Repeat the scan until NO orphan messages remain."),
    ("5", "The resulting frontier is the RECOVERY LINE — a consistent global state."),
]

for i, (num, desc) in enumerate(steps):
    y = 2.4 + i * 0.95
    add_rounded_box(slide, 0.7, y, 0.55, 0.6, GREEN, num, 16, BG)
    add_textbox(slide, 1.4, y + 0.05, 11.3, 0.6, desc, size=15, color=WHITE)

# ============================================================
# SLIDE 6: SIMULATOR FEATURES
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BG)

add_accent_bar(slide, 0.7, 0.6, 0.9, ACCENT, 0.12)
add_textbox(slide, 0.95, 0.55, 11, 0.9, "Simulator Features", size=36, color=ACCENT, bold=True)

features = [
    ("Process Management", "Add / remove distributed process timelines.", ACCENT),
    ("Checkpointing", "Take local checkpoints on any process.", ORANGE),
    ("Message Passing", "Send messages between processes with send/receive events.", TEAL),
    ("Failure Simulation", "Mark any process as failed at any time.", RED),
    ("Recovery Line", "One-click computation of the consistent global state.", GREEN),
    ("Two Demos", "Domino Effect (uncoordinated) & Coordinated Recovery.", ACCENT),
    ("Event Log", "Color-coded log narrating every rollback step.", MUTED),
]

cols = 2
for i, (title, desc, color) in enumerate(features):
    col = i % cols
    row = i // cols
    x = 0.7 + col * 6.1
    y = 1.7 + row * 1.65
    add_rounded_box(slide, x, y, 5.8, 1.35, BG2, "", 12)
    add_textbox(slide, x + 0.2, y + 0.12, 5.4, 0.4, title, size=15, color=color, bold=True)
    add_textbox(slide, x + 0.2, y + 0.55, 5.4, 0.7, desc, size=12, color=MUTED)

# ============================================================
# SLIDE 7: DEMO 1 — DOMINO EFFECT
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BG)

add_accent_bar(slide, 0.7, 0.6, 0.9, RED, 0.12)
add_textbox(slide, 0.95, 0.55, 11, 0.9, "Demo 1: The Domino Effect", size=36, color=RED, bold=True)

add_textbox(slide, 0.7, 1.6, 12, 0.5, "Scenario: 3 processes (P0, P1, P2) with UNCOORDINATED checkpoints forming a cycle.", size=16, color=WHITE)

add_rounded_box(slide, 0.7, 2.2, 12.0, 0.7, BG2,
    "P0 → P1 (m1)  →  P1 → P2 (m2)  →  P2 → P0 (m3)   [checkpoints taken independently]",
    16, ORANGE)

add_textbox(slide, 0.7, 3.1, 12, 0.5, "When P2 fails, the recovery triggers a CASCADE:", size=16, color=ACCENT, bold=True)

cascade = [
    ("P2 rolls back", "orphan message m2 — its send (from P1) is not in the cut."),
    ("P0 rolls back", "orphan message m3 — its send (from P2) is not in the cut."),
    ("P1 rolls back", "orphan message m1 — its send (from P0) is not in the cut."),
]

for i, (title, desc) in enumerate(cascade):
    y = 3.7 + i * 0.85
    add_rounded_box(slide, 0.7, y, 2.5, 0.6, RED, title, 14, WHITE)
    add_textbox(slide, 3.4, y + 0.08, 9.3, 0.5, desc, size=14, color=WHITE)

add_rounded_box(slide, 0.7, 6.4, 12.0, 0.7, RED,
    "Result: All processes roll back to the very beginning — total loss of work!", 16, WHITE)

# ============================================================
# SLIDE 8: DEMO 2 — COORDINATED RECOVERY
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BG)

add_accent_bar(slide, 0.7, 0.6, 0.9, GREEN, 0.12)
add_textbox(slide, 0.95, 0.55, 11, 0.9, "Demo 2: Coordinated Recovery", size=36, color=GREEN, bold=True)

add_textbox(slide, 0.7, 1.6, 12, 0.5, "Scenario: 2 processes that checkpoint CONSISTENTLY (both checkpoint before & after messaging).", size=16, color=WHITE)

add_rounded_box(slide, 0.7, 2.2, 12.0, 0.7, BG2,
    "[P0 cp] [P1 cp]  →  P0↔P1 messages  →  [P0 cp] [P1 cp]",
    16, ORANGE)

add_bullets(slide, 0.7, 3.2, 12, 3.0, [
    "Both processes take checkpoints at coordinated points.",
    "No message crosses a checkpoint boundary inconsistently.",
    "When P1 fails, the recovery line check finds ZERO orphan messages.",
    "No rollbacks needed — recovery is instant and minimal.",
], size=17)

add_rounded_box(slide, 0.7, 6.2, 12.0, 0.9, GREEN,
    "Result: 0 rollbacks — Coordinated checkpointing eliminates the Domino Effect!", 18, WHITE)

# ============================================================
# SLIDE 9: TECHNOLOGY STACK
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BG)

add_accent_bar(slide, 0.7, 0.6, 0.9, ACCENT, 0.12)
add_textbox(slide, 0.95, 0.55, 11, 0.9, "Technology Stack", size=36, color=ACCENT, bold=True)

stack = [
    ("HTML5", "Page structure & layout", ORANGE),
    ("CSS3", "Dark-themed responsive styling", TEAL),
    ("JavaScript (ES6)", "Simulation engine, recovery algorithm, Canvas rendering", ACCENT),
    ("HTML5 Canvas", "Visual rendering of process timelines, events, messages & recovery line", RED),
]

for i, (tech, desc, color) in enumerate(stack):
    y = 1.8 + i * 1.2
    add_rounded_box(slide, 0.7, y, 3.0, 0.9, color, tech, 20, BG)
    add_textbox(slide, 4.0, y + 0.2, 8.7, 0.6, desc, size=16, color=WHITE)

add_textbox(slide, 0.7, 6.7, 12, 0.5, "No external libraries or frameworks — pure vanilla implementation.", size=14, color=MUTED)

# ============================================================
# SLIDE 10: CONCLUSION & FUTURE WORK
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BG)

add_accent_bar(slide, 0.7, 0.6, 0.9, ACCENT, 0.12)
add_textbox(slide, 0.95, 0.55, 11, 0.9, "Conclusion & Future Work", size=36, color=ACCENT, bold=True)

add_textbox(slide, 0.7, 1.7, 12, 0.5, "Conclusion:", size=18, color=GREEN, bold=True)
add_bullets(slide, 0.7, 2.2, 12, 1.8, [
    "The simulator visually demonstrates the core challenge of distributed rollback recovery.",
    "It clearly contrasts uncoordinated (domino effect) vs coordinated (minimal rollback) checkpointing.",
    "The recovery line algorithm correctly finds a consistent global state using orphan-message elimination.",
], size=15)

add_textbox(slide, 0.7, 4.2, 12, 0.5, "Future Work:", size=18, color=ORANGE, bold=True)
add_bullets(slide, 0.7, 4.7, 12, 2.3, [
    "Implement message logging (sender-based / receiver-based) for optimistic recovery.",
    "Add step-by-step animation of each rollback iteration.",
    "Support for concurrent checkpointing algorithms (Chandy-Lamport snapshot).",
    "Export / import simulation scenarios.",
], size=15)

# ============================================================
# SLIDE 11: THANK YOU
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, BG)

add_textbox(slide, 1, 2.8, 11.3, 1.2, "Thank You", size=56, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
add_textbox(slide, 1, 4.1, 11.3, 0.6, "Distributed Rollback Recovery Simulator", size=22, color=MUTED, align=PP_ALIGN.CENTER)
add_textbox(slide, 1, 4.9, 11.3, 0.5, "Questions & Discussion", size=16, color=TEAL, align=PP_ALIGN.CENTER)

output = r"d:\MCA SEM 1\Distributed Computing\PROJECT TINY\Distributed_Rollback_Recovery_Simulator.pptx"
prs.save(output)
print(f"PPT saved to: {output}")
